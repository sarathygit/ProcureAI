"""
generate_samples.py — Deterministic sample data generator for ProcureAI
=======================================================================

Creates inside ./samples :

  Quotations
  1. supplier_alpha_compliant.pdf     — Apex Solutions: fully compliant, strong track record
  2. supplier_beta_noncompliant.pdf   — Beta Tech Enterprises: cheaper but fails RAM / storage / warranty,
                                        and its GSTIN fails the checksum
  3. supplier_gamma_incomplete.xlsx   — Gamma Systems: Warranty & Delivery blank -> NEEDS_CLARIFICATION (MSME)
  4. supplier_delta_messy.pdf         — Delta Infotech: messy wording ("half a terabyte", "a fortnight"),
                                        hidden costs in the fine print and a suspiciously low price
  5. supplier_epsilon_clone.pdf       — Epsilon Traders: a near-copy of Apex's quotation (bid-collusion signal)

  Reference data (simulated)
  6. supplier_history.json            — delivery track record & capacity per supplier
  7. procurement_history.json         — 12 months of past procurement events for the analytics dashboard

Run:  python generate_samples.py
"""

from __future__ import annotations

import json
import os
import random
from datetime import date, timedelta
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

SAMPLES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "samples")

GSTIN_CHARS = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"


def gstin(first14: str, valid: bool = True) -> str:
    total = 0
    for i, ch in enumerate(first14):
        v = GSTIN_CHARS.index(ch) * (1 if i % 2 == 0 else 2)
        total += v // 36 + v % 36
    check = GSTIN_CHARS[(36 - total % 36) % 36]
    if not valid:
        check = GSTIN_CHARS[(GSTIN_CHARS.index(check) + 7) % 36]
    return first14 + check


# Helvetica (the ReportLab default) has no rupee glyph, so PDFs use the "INR" prefix.

APEX_TEMPLATE = {
    "intro": ("We are pleased to offer the <b>{model}</b> laptop in the following configuration: "
              "{ram} GB RAM, {storage} GB SSD, backed by {warranty} Months Warranty (on-site). Standard lead time is "
              "{delivery} Days Delivery from receipt of purchase order."),
    "section_offer": "Commercial Offer",
    "section_table": "Technical &amp; Commercial Summary",
    "labels": ["Parameter", "Offered Value", "Model", "Memory", "Storage", "Warranty Period", "Lead Time",
               "Rate per unit (INR)", "GST", "Freight (INR, total)"],
    "terms": [
        "1. Payment terms: 100% within 30 days of delivery and invoice.",
        "2. Prices are firm and valid for 30 days from the date of issue.",
        "3. Delivery to a single consignee location; installation support included.",
        "4. Any statutory levies introduced after the quotation date will be charged extra.",
    ],
    "phone": "Tel: +91 80 4123 9876 | sales@apexsolutions.example",
}

BETA_TEMPLATE = {
    "intro": ("Further to your enquiry, Beta Tech Enterprises is happy to propose the <b>{model}</b>. Configuration: "
              "{ram} GB RAM, {storage} GB SSD. All units carry {warranty} Months Warranty through our carry-in service "
              "centres in Delhi NCR. We can complete supply in {delivery} Days Delivery after PO confirmation."),
    "section_offer": "Our Proposal",
    "section_table": "Specification Sheet",
    "labels": ["Item", "Details", "Product", "Installed memory", "Disk", "Service cover", "Supply schedule",
               "Price each (INR)", "Tax", "Transport (INR)"],
    "terms": [
        "a) 50% advance with PO, balance against delivery.",
        "b) Quotation validity: 21 days.",
        "c) Goods once sold will not be taken back. Disputes subject to Delhi jurisdiction.",
    ],
    "phone": "Phone: 011-4567 2210 | betatech.delhi@example.com",
}

SAMPLE_PDF_QUOTES: List[Dict] = [
    {
        "filename": "supplier_alpha_compliant.pdf", "title": "Apex Solutions Quotation", "supplier": "Apex Solutions",
        "address": "Plot 14, Electronic City Phase II, Bengaluru 560100",
        "gstin": gstin("29ABCDE1234F1Z"), "quote_ref": "APX/Q/2026/1187",
        "unit_price": 52000, "ram": 16, "storage": 512, "warranty": 24, "delivery": 7, "gst": 18, "shipping": 2000,
        "model": "Apex ProBook 14 — Intel Core i7, 14\" FHD IPS", "accent": colors.HexColor("#1F4E79"),
        "template": APEX_TEMPLATE,
    },
    {
        "filename": "supplier_beta_noncompliant.pdf", "title": "Beta Tech Enterprises", "supplier": "Beta Tech Enterprises",
        "address": "B-42, Okhla Industrial Area Phase I, New Delhi 110020",
        "gstin": gstin("07FGHIJ5678K1Z", valid=False), "quote_ref": "BTE-2026-0942",
        "unit_price": 41000, "ram": 8, "storage": 256, "warranty": 12, "delivery": 10, "gst": 18, "shipping": 1500,
        "model": "BetaBook Lite 15 — Intel Core i5, 15.6\" HD", "accent": colors.HexColor("#7B2C2C"),
        "template": BETA_TEMPLATE,
    },
    {
        "filename": "supplier_epsilon_clone.pdf", "title": "Epsilon Traders Quotation", "supplier": "Epsilon Traders",
        "address": "Plot 41, Electronic City Phase II, Bengaluru 560100",
        "gstin": gstin("29KLMNO4321P1Z"), "quote_ref": "EPS/Q/2026/1188",
        "unit_price": 53000, "ram": 16, "storage": 512, "warranty": 24, "delivery": 8, "gst": 18, "shipping": 2000,
        "model": "Apex ProBook 14 — Intel Core i7, 14\" FHD IPS", "accent": colors.HexColor("#1F4E79"),
        "template": APEX_TEMPLATE,   # deliberately re-uses Apex's wording and phone number
    },
]


def _fmt(n: float) -> str:
    """Indian digit grouping without currency symbol (e.g. 10,40,000)."""
    whole = int(round(n))
    s = str(whole)
    if len(s) <= 3:
        return s
    head, tail = s[:-3], s[-3:]
    groups = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    if head:
        groups.insert(0, head)
    return ",".join(groups) + "," + tail


def build_pdf_quotation(spec: Dict, out_dir: str) -> str:
    path = os.path.join(out_dir, spec["filename"])
    tpl = spec["template"]
    doc = SimpleDocTemplate(
        path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
        title=spec["title"], author=spec["supplier"],
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("QTitle", parent=styles["Title"], textColor=spec["accent"], fontSize=20, spaceAfter=4)
    sub_style = ParagraphStyle("QSub", parent=styles["Normal"], fontSize=8.5, textColor=colors.grey, alignment=TA_CENTER)
    h_style = ParagraphStyle("QH", parent=styles["Heading3"], textColor=spec["accent"], spaceBefore=8, spaceAfter=4)
    body = ParagraphStyle("QBody", parent=styles["Normal"], fontSize=10, leading=14)
    small = ParagraphStyle("QSmall", parent=styles["Normal"], fontSize=8.5, leading=11, textColor=colors.darkgrey)

    qty_example = 20
    base = spec["unit_price"] * qty_example
    gst_amt = base * spec["gst"] / 100.0
    total = base + gst_amt + spec["shipping"]
    L = tpl["labels"]

    story = [
        Paragraph(spec["title"], title_style),
        Paragraph(f"{spec['address']} | GSTIN {spec['gstin']}", sub_style),
        Paragraph(tpl["phone"], sub_style),
        Spacer(1, 8),
        Paragraph(f"<b>Quotation Ref:</b> {spec['quote_ref']} &nbsp;&nbsp; <b>Date:</b> 20 Oct 2026", body),
        Paragraph("<b>To:</b> Procurement Department, Corporate IT Refresh Programme", body),
        Spacer(1, 6),
        Paragraph(tpl["section_offer"], h_style),
        Paragraph(tpl["intro"].format(model=spec["model"], ram=spec["ram"], storage=spec["storage"],
                                      warranty=spec["warranty"], delivery=spec["delivery"]), body),
        Spacer(1, 6),
        Paragraph(f"Unit Price: INR {_fmt(spec['unit_price'])} per laptop (exclusive of GST).", body),
        Paragraph(f"GST: {spec['gst']}% applicable on base value.", body),
        Paragraph(f"Shipping: INR {_fmt(spec['shipping'])} (total, for the complete consignment).", body),
        Spacer(1, 8),
        Paragraph(tpl["section_table"], h_style),
    ]

    table_data = [
        [L[0], L[1]],
        [L[2], spec["model"]],
        [L[3], f"{spec['ram']} GB DDR5"],
        [L[4], f"{spec['storage']} GB NVMe"],
        [L[5], f"{spec['warranty']} months"],
        [L[6], f"{spec['delivery']} calendar days"],
        [L[7], _fmt(spec["unit_price"])],
        [L[8], f"{spec['gst']}%"],
        [L[9], _fmt(spec["shipping"])],
    ]
    table = Table(table_data, colWidths=[60 * mm, 110 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), spec["accent"]),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.lightgrey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(table)
    story.append(Spacer(1, 8))
    story.append(Paragraph("Indicative Value for 20 Units", h_style))

    value_table = Table([
        ["Base value (20 x INR " + _fmt(spec["unit_price"]) + ")", "INR " + _fmt(base)],
        [f"GST @ {spec['gst']}%", "INR " + _fmt(gst_amt)],
        ["Freight", "INR " + _fmt(spec["shipping"])],
        ["Indicative landed value", "INR " + _fmt(total)],
    ], colWidths=[110 * mm, 60 * mm])
    value_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("LINEABOVE", (0, -1), (-1, -1), 0.8, spec["accent"]),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(value_table)
    story.append(Spacer(1, 10))
    story.append(Paragraph("Terms &amp; Conditions", h_style))
    for line in tpl["terms"]:
        story.append(Paragraph(line, small))
    story.append(Spacer(1, 16))
    story.append(Paragraph(f"For <b>{spec['supplier']}</b><br/><br/>Authorised Signatory", body))

    doc.build(story)
    return path


def build_messy_pdf(out_dir: str) -> str:
    """A realistic, badly-structured quotation: no tables, informal units, costs hidden in the fine print."""
    path = os.path.join(out_dir, "supplier_delta_messy.pdf")
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title="Delta Infotech offer", author="Delta Infotech")
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontSize=10.5, leading=15)
    tiny = ParagraphStyle("t", parent=ss["Normal"], fontSize=7, leading=9, textColor=colors.grey)
    head = ParagraphStyle("h", parent=ss["Heading2"], textColor=colors.HexColor("#4a148c"))
    story = [
        Paragraph("Delta Infotech", head),
        Paragraph(f"No. 7, Anna Salai, Chennai 600002 · GSTIN {gstin('33PQRST9876U1Z')} · Udyam Regn UDYAM-TN-02-0034567",
                  tiny),
        Spacer(1, 10),
        Paragraph("Dear Sir/Madam,", body),
        Paragraph("Greetings from Delta Infotech! With reference to your requirement for business laptops, please find our "
                  "best offer below. We have supplied to many reputed corporates across South India.", body),
        Spacer(1, 6),
        Paragraph("<b>What you get:</b> Delta Edge 14 notebook, Ryzen 5 processor, memory of 16 gigs (DDR4), "
                  "half a terabyte NVMe drive, 14 inch full-HD display, backlit keyboard.", body),
        Paragraph("<b>Cover:</b> 2 yrs onsite cover from the date of invoice.", body),
        Paragraph("<b>Timeline:</b> we will dispatch within a fortnight of receiving your PO.", body),
        Paragraph("<b>Commercials:</b> per-unit rate Rs 36,500/- plus GST @18.", body),
        Spacer(1, 8),
        Paragraph("Looking forward to a long association.", body),
        Paragraph("Warm regards,<br/>K. Ramesh, Business Development", body),
        Spacer(1, 30),
        Paragraph("Notes: (1) Freight extra at actuals. (2) Installation and OS imaging charged extra at Rs 750 per unit. "
                  "(3) Battery covered for 6 months only. (4) Prices valid for 7 days, subject to stock availability. "
                  "(5) Payment within 15 days.", tiny),
    ]
    doc.build(story)
    return path


def build_excel_quotation(out_dir: str) -> str:
    path = os.path.join(out_dir, "supplier_gamma_incomplete.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"

    header_fill = PatternFill("solid", fgColor="2E7D32")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    rows = [
        ("Parameter", "Value"),
        ("Supplier Name", "Gamma Systems"),
        ("Quotation Reference", "GS/2026/Q-311"),
        ("Product", "Gamma Edge 14 — AMD Ryzen 7, 14\" FHD"),
        ("Unit Price (INR)", 48000),
        ("RAM (GB)", 16),
        ("Storage (GB SSD)", 512),
        ("Warranty (Months)", None),   # intentionally blank
        ("Delivery Days", None),       # intentionally blank
        ("GST (%)", 18),
        ("Shipping Cost (INR)", 1800),
        ("GSTIN", gstin("27UVWXY2468Z1Z")),
        ("MSME / Udyam Registered", "Yes - UDYAM-MH-26-0012345"),
        ("Quote Validity (Days)", 30),
        ("Payment Terms (Days)", 45),
    ]
    for r_idx, (param, value) in enumerate(rows, start=1):
        c1 = ws.cell(row=r_idx, column=1, value=param)
        c2 = ws.cell(row=r_idx, column=2, value=value)
        for c in (c1, c2):
            c.border = border
            c.alignment = Alignment(vertical="center")
        if r_idx == 1:
            for c in (c1, c2):
                c.fill = header_fill
                c.font = header_font
        else:
            c1.font = Font(bold=True)

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 42
    wb.save(path)
    return path


SUPPLIER_HISTORY = {
    "Apex Solutions": {"orders_completed": 42, "on_time_rate": 0.93, "defect_rate": 0.012, "disputes": 0,
                       "years_active": 9, "avg_delay_days": 1.0, "capacity_units": 12},
    "Beta Tech Enterprises": {"orders_completed": 25, "on_time_rate": 0.70, "defect_rate": 0.06, "disputes": 3,
                              "years_active": 6, "avg_delay_days": 6.5, "capacity_units": 100},
    "Gamma Systems": {"orders_completed": 18, "on_time_rate": 0.89, "defect_rate": 0.02, "disputes": 0,
                      "years_active": 4, "avg_delay_days": 2.0, "capacity_units": 40},
    "Delta Infotech": {"orders_completed": 3, "on_time_rate": 0.55, "defect_rate": 0.08, "disputes": 2,
                       "years_active": 1, "avg_delay_days": 9.0, "capacity_units": 50},
    "Old Mart Computers": {"orders_completed": 12, "on_time_rate": 0.60, "defect_rate": 0.10, "disputes": 6,
                           "years_active": 8, "avg_delay_days": 12.0, "capacity_units": 30, "blacklisted": True},
}


def build_procurement_history(out_dir: str) -> str:
    rng = random.Random(42)
    categories = [("Laptops", 45000, 10, 40), ("Monitors", 11000, 10, 60), ("Servers", 380000, 1, 4),
                  ("Networking", 26000, 4, 20), ("Peripherals", 1800, 20, 150)]
    suppliers = ["Apex Solutions", "Gamma Systems", "Beta Tech Enterprises", "Delta Infotech", "Northwind IT",
                 "Contoso Systems"]
    start = date(2025, 11, 1)
    events = []
    for i in range(48):
        d = start + timedelta(days=int(i * 7.4) + rng.randint(0, 4))
        cat, price, qmin, qmax = rng.choice(categories)
        qty = rng.randint(qmin, qmax)
        n_quotes = rng.randint(2, 6)
        quotes = sorted(price * rng.uniform(0.82, 1.25) for _ in range(n_quotes))
        compliant = max(1, n_quotes - rng.randint(0, 2))
        lowest_any = quotes[0]
        winner_price = quotes[min(rng.randint(0, 1), len(quotes) - 1)]
        awarded = winner_price * qty * 1.18
        events.append({
            "date": d.isoformat(), "category": cat, "quantity": qty, "quotes_received": n_quotes,
            "compliant_quotes": compliant, "winner": rng.choice(suppliers),
            "awarded_value": round(awarded, 2),
            "highest_quote_value": round(quotes[-1] * qty * 1.18, 2),
            "lowest_quote_value": round(lowest_any * qty * 1.18, 2),
            "negotiated_savings": round(awarded * rng.uniform(0.0, 0.06), 2),
            "cycle_days_manual_baseline": rng.randint(9, 21),
            "cycle_days": rng.randint(2, 6),
            "clarifications": rng.randint(0, 3),
            "on_time": rng.random() < 0.86,
        })
    path = os.path.join(out_dir, "procurement_history.json")
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(events, fh, indent=1)
    return path


def generate_all(out_dir: str = SAMPLES_DIR) -> List[str]:
    os.makedirs(out_dir, exist_ok=True)
    paths = [build_pdf_quotation(spec, out_dir) for spec in SAMPLE_PDF_QUOTES]
    paths.append(build_excel_quotation(out_dir))
    paths.append(build_messy_pdf(out_dir))
    hist = os.path.join(out_dir, "supplier_history.json")
    with open(hist, "w", encoding="utf-8") as fh:
        json.dump(SUPPLIER_HISTORY, fh, indent=1)
    paths.append(hist)
    paths.append(build_procurement_history(out_dir))
    return paths


QUOTATION_FILES = [
    "supplier_alpha_compliant.pdf", "supplier_beta_noncompliant.pdf", "supplier_gamma_incomplete.xlsx",
    "supplier_delta_messy.pdf", "supplier_epsilon_clone.pdf",
]

if __name__ == "__main__":
    for p in generate_all():
        print(f"Created {os.path.relpath(p)}")
