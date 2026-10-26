"""
extractor.py — Quotation ingestion for ProcureAI
================================================

Turns supplier quotation files into structured `SupplierQuotation` models.

  * PDF   -> PyMuPDF (fitz) full-text extraction + deterministic regex rules
  * Excel -> openpyxl cell scan with keyword matching on "Parameter / Value" style sheets

Golden rule: if a mandatory value cannot be found it stays `None`. The extractor never
guesses or fills defaults for mandatory specifications — the decision engine turns those
gaps into NEEDS_CLARIFICATION.
"""

from __future__ import annotations

import io
import os
import re
from typing import Any, Dict, List, Optional, Tuple

import warnings

with warnings.catch_warnings():
    warnings.simplefilter("ignore", FutureWarning)  # newer PyMuPDF prefers `import pymupdf`; `fitz` remains supported
    import fitz  # PyMuPDF
from openpyxl import load_workbook

from engine import SupplierQuotation

# ---------------------------------------------------------------------------
# Regex rules (primary patterns from the spec, plus key/value fallbacks)
# ---------------------------------------------------------------------------

FLAGS = re.IGNORECASE
CURRENCY = r"(?:₹|\$|Rs\.?|INR)?"

RAM_PATTERNS = [
    re.compile(r"(\d+)\s*(?:GB)?\s*RAM", FLAGS),
    re.compile(r"(?:RAM|Memory)\s*(?:\(GB\))?[:\s\-]*(\d+)\s*GB", FLAGS),
]
STORAGE_PATTERNS = [
    re.compile(r"(\d+)\s*(GB|TB)\s*(?:SSD|Storage)", FLAGS),
    re.compile(r"(?:Storage|SSD)\s*(?:\(GB\))?[:\s\-]*(\d+)\s*(GB|TB)", FLAGS),
]
PRICE_PATTERNS = [
    re.compile(r"Unit\s*Price[:\s]*" + CURRENCY + r"\s*([0-9,]+(?:\.[0-9]{2})?)", FLAGS),
    re.compile(r"(?:Price|Cost|Rate|Amount)[:\s]*" + CURRENCY + r"\s*([0-9,]+(?:\.[0-9]{2})?)", FLAGS),
]
WARRANTY_PATTERNS = [
    re.compile(r"(\d+)\s*(Months?|Years?)\s*Warranty", FLAGS),
    re.compile(r"Warranty(?:\s*Period)?[:\s\-]*(\d+)\s*(Months?|Years?)", FLAGS),
]
DELIVERY_PATTERNS = [
    re.compile(r"(\d+)\s*Days?\s*Delivery", FLAGS),
    re.compile(r"(?:Delivery|Lead\s*Time)(?:\s*Days)?[:\s\-]*(?:within\s*)?(\d+)\s*(?:calendar\s*|working\s*)?Days?", FLAGS),
]
SHIPPING_PATTERNS = [
    re.compile(r"(?:Shipping|Freight)[:\s]*" + CURRENCY + r"\s*([0-9,]+)", FLAGS),
    re.compile(r"(?:Shipping|Freight)\s*(?:\(INR[^)]*\)|Cost|Charges)?[:\s]*" + CURRENCY + r"\s*([0-9,]+)", FLAGS),
]
GST_PATTERNS = [
    re.compile(r"GST[^0-9\n]{0,20}(\d+(?:\.\d+)?)\s*%", FLAGS),
    re.compile(r"(\d+(?:\.\d+)?)\s*%\s*GST", FLAGS),
]
SUPPLIER_PATTERNS = [
    re.compile(r"(?:Supplier|Vendor|Company)\s*(?:Name)?\s*[:\-]\s*(.+)", FLAGS),
    re.compile(r"For\s+(.+?)\s*\n\s*Authori[sz]ed\s+Signatory", FLAGS),
]


def _clean_number(text: str) -> Optional[float]:
    try:
        return float(text.replace(",", "").strip())
    except (ValueError, AttributeError):
        return None


def _first_match(patterns: List[re.Pattern], text: str) -> Optional[re.Match]:
    for pattern in patterns:
        m = pattern.search(text)
        if m:
            return m
    return None


def _supplier_from_filename(filename: str) -> str:
    stem = os.path.splitext(os.path.basename(filename))[0]
    stem = re.sub(r"(?i)^(supplier|vendor|quote|quotation)[_\-\s]*", "", stem)
    stem = re.sub(r"(?i)[_\-\s]*(compliant|noncompliant|incomplete|quotation|quote)$", "", stem)
    words = [w for w in re.split(r"[_\-\s]+", stem) if w]
    return " ".join(w.capitalize() for w in words) or "Unknown Supplier"


def _clean_supplier_name(raw: str) -> str:
    name = raw.strip().splitlines()[0].strip()
    name = re.sub(r"(?i)\s*[-–—:]?\s*(?:quotation|quote|commercial offer)\s*$", "", name).strip()
    return name


# ---------------------------------------------------------------------------
# Text -> quotation (shared by PDF and any plain-text source)
# ---------------------------------------------------------------------------


def parse_quotation_text(text: str, filename: str) -> SupplierQuotation:
    # Supplier name: explicit label, then signatory block, then first non-empty line, then filename.
    supplier: Optional[str] = None
    m = _first_match(SUPPLIER_PATTERNS, text)
    if m:
        supplier = _clean_supplier_name(m.group(1))
    if not supplier:
        first_line = next((ln.strip() for ln in text.splitlines() if ln.strip()), "")
        if first_line and len(first_line) <= 80 and not re.search(r"\d{3,}", first_line):
            supplier = _clean_supplier_name(first_line)
    if not supplier:
        supplier = _supplier_from_filename(filename)

    ram = None
    m = _first_match(RAM_PATTERNS, text)
    if m:
        ram = int(m.group(1))

    storage = None
    m = _first_match(STORAGE_PATTERNS, text)
    if m:
        storage = int(m.group(1))
        if m.group(2).upper() == "TB":
            storage *= 1024

    price = None
    m = _first_match(PRICE_PATTERNS, text)
    if m:
        price = _clean_number(m.group(1))

    warranty = None
    m = _first_match(WARRANTY_PATTERNS, text)
    if m:
        warranty = int(m.group(1))
        if m.group(2).lower().startswith("year"):
            warranty *= 12

    delivery = None
    m = _first_match(DELIVERY_PATTERNS, text)
    if m:
        delivery = int(m.group(1))

    shipping = 0.0
    m = _first_match(SHIPPING_PATTERNS, text)
    if m:
        shipping = _clean_number(m.group(1)) or 0.0

    gst = 18.0
    m = _first_match(GST_PATTERNS, text)
    if m:
        value = _clean_number(m.group(1))
        if value is not None:
            gst = value

    return SupplierQuotation(
        supplier_name=supplier,
        unit_price=price,
        ram_gb=ram,
        storage_gb=storage,
        warranty_months=warranty,
        delivery_days=delivery,
        gst_percent=gst,
        shipping_cost_total=shipping,
    )


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def extract_pdf_text(file_bytes: bytes) -> str:
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        return "\n".join(page.get_text("text") for page in doc)


def parse_pdf_quotation(file_bytes: bytes, filename: str) -> SupplierQuotation:
    text = extract_pdf_text(file_bytes)
    return parse_quotation_text(text, filename)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

# Order matters: more specific keys are tested first (e.g. "shipping cost" before "cost").
EXCEL_KEYWORDS: List[Tuple[str, List[str]]] = [
    ("supplier_name", ["supplier name", "vendor name", "company name", "supplier", "vendor", "company"]),
    ("shipping_cost_total", ["shipping", "freight", "delivery charge", "transport"]),
    ("gst_percent", ["gst", "tax"]),
    ("delivery_days", ["delivery days", "delivery", "lead time", "dispatch"]),
    ("warranty_months", ["warranty"]),
    ("ram_gb", ["ram", "memory"]),
    ("storage_gb", ["storage", "ssd", "hdd", "disk"]),
    ("unit_price", ["unit price", "price", "rate", "unit cost", "cost"]),
]


def _match_excel_key(label: str) -> Optional[str]:
    norm = re.sub(r"\s+", " ", label.strip().lower())
    if not norm or norm in {"parameter", "value", "field", "description"}:
        return None
    for field, keywords in EXCEL_KEYWORDS:
        for kw in keywords:
            if re.search(r"(?<![a-z])" + re.escape(kw) + r"(?![a-z])", norm):
                return field
    return None


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() in {"", "-", "—", "n/a", "na", "tbd", "tba"})


def _excel_number(value: Any) -> Optional[float]:
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    m = re.search(r"-?[0-9][0-9,]*(?:\.[0-9]+)?", str(value))
    return _clean_number(m.group(0)) if m else None


def _convert_excel_value(field: str, label: str, value: Any) -> Any:
    if field == "supplier_name":
        return None if _is_blank(value) else str(value).strip()
    number = _excel_number(value)
    if number is None:
        return None
    text = f"{label} {value}".lower()
    if field == "storage_gb":
        return int(number * 1024) if re.search(r"\btb\b", text) else int(number)
    if field == "warranty_months":
        return int(number * 12) if re.search(r"\b(?:years?|yrs?)\b", text) else int(number)
    if field == "delivery_days":
        return int(number * 7) if re.search(r"\bweeks?\b", text) else int(number)
    if field == "ram_gb":
        return int(number)
    return float(number)


def parse_excel_quotation(file_bytes: bytes, filename: str) -> SupplierQuotation:
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    found: Dict[str, Any] = {}
    seen: set = set()

    for ws in wb.worksheets:
        # Layout 1: vertical key/value pairs (label in a cell, value in the next cell to the right).
        for row in ws.iter_rows(values_only=True):
            cells = list(row)
            # The label is the first text cell of the row; later text cells are values, not keys.
            idx = next((i for i, c in enumerate(cells) if isinstance(c, str) and c.strip()), None)
            if idx is None:
                continue
            label = cells[idx]
            field = _match_excel_key(label)
            if field is None or field in seen:
                continue
            value = cells[idx + 1] if idx + 1 < len(cells) else None
            seen.add(field)  # a labelled-but-blank field stays None — it is NOT searched elsewhere
            converted = _convert_excel_value(field, label, value)
            if converted is not None:
                found[field] = converted

        # Layout 2: horizontal header row with one data row beneath it.
        if not seen:
            rows = list(ws.iter_rows(values_only=True))
            for r_idx, row in enumerate(rows[:-1]):
                header_fields = {i: _match_excel_key(c) for i, c in enumerate(row) if isinstance(c, str)}
                header_fields = {i: f for i, f in header_fields.items() if f}
                if len(header_fields) >= 3:
                    data_row = rows[r_idx + 1]
                    for i, field in header_fields.items():
                        if field in seen:
                            continue
                        seen.add(field)
                        value = data_row[i] if i < len(data_row) else None
                        converted = _convert_excel_value(field, str(row[i]), value)
                        if converted is not None:
                            found[field] = converted
                    break

    supplier = found.get("supplier_name") or _supplier_from_filename(filename)
    return SupplierQuotation(
        supplier_name=supplier,
        unit_price=found.get("unit_price"),
        ram_gb=found.get("ram_gb"),
        storage_gb=found.get("storage_gb"),
        warranty_months=found.get("warranty_months"),
        delivery_days=found.get("delivery_days"),
        gst_percent=found.get("gst_percent", 18.0),
        shipping_cost_total=found.get("shipping_cost_total", 0.0),
    )


def extract_excel_text(file_bytes: bytes) -> str:
    """Flatten a workbook into readable text (used for the raw-extraction preview in the UI)."""
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    lines: List[str] = []
    for ws in wb.worksheets:
        lines.append(f"[Sheet: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                lines.append(" | ".join("" if c is None else str(c) for c in row))
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Dispatcher
# ---------------------------------------------------------------------------


def parse_quotation_file(file_bytes: bytes, filename: str) -> Tuple[SupplierQuotation, str]:
    """Parse any supported file. Returns (quotation, raw_text_preview)."""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        text = extract_pdf_text(file_bytes)
        return parse_quotation_text(text, filename), text
    if ext in {".xlsx", ".xlsm"}:
        return parse_excel_quotation(file_bytes, filename), extract_excel_text(file_bytes)
    raise ValueError(f"Unsupported file type '{ext}'. Please upload .pdf or .xlsx quotations.")


if __name__ == "__main__":
    import sys

    targets = sys.argv[1:]
    if not targets and os.path.isdir("samples"):
        targets = [os.path.join("samples", f) for f in sorted(os.listdir("samples"))]
    for path in targets:
        with open(path, "rb") as fh:
            quotation, _ = parse_quotation_file(fh.read(), os.path.basename(path))
        print(f"{path}: {quotation.model_dump()}")
