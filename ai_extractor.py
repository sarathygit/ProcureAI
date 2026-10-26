"""
ai_extractor.py — Hybrid AI extraction with per-field confidence (Feature #2)
=============================================================================

Every quotation goes through two independent readers:

  1. Deterministic rules  — the regex / cell rules from extractor.py, plus low-confidence
                            "heuristic" rules for messy wording ("2 yrs cover", "a fortnight").
  2. Claude (optional)    — reads the text (or the raw PDF for scanned documents) and returns
                            a value, a confidence and the exact evidence sentence per field.

The two are reconciled field by field. Agreement raises confidence, disagreement lowers it,
and anything below REVIEW_THRESHOLD must be confirmed by a human before the decision engine
will use it. Low-confidence values are never silently trusted.
"""

from __future__ import annotations

import io
import os
import re
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from pydantic import BaseModel, Field

import extractor as rules
import llm
from engine import FIELD_LABELS, MANDATORY_FIELDS, SupplierQuotation, normalise_casual_text

REVIEW_THRESHOLD = 0.80

NUMERIC_FIELDS = ["unit_price", "ram_gb", "storage_gb", "warranty_months", "delivery_days",
                  "gst_percent", "shipping_cost_total", "quote_validity_days", "payment_terms_days"]
INT_FIELDS = {"ram_gb", "storage_gb", "warranty_months", "delivery_days", "quote_validity_days",
              "payment_terms_days"}

EXTRA_LABELS = {
    "supplier_name": "Supplier Name",
    "gstin": "GSTIN",
    "is_msme": "MSME / Udyam registered",
    "quote_validity_days": "Quote Validity (Days)",
    "payment_terms_days": "Payment Terms (Days)",
}


def field_label(field: str) -> str:
    return FIELD_LABELS.get(field) or EXTRA_LABELS.get(field, field)


# ---------------------------------------------------------------------------
# Result models
# ---------------------------------------------------------------------------


class FieldResult(BaseModel):
    field: str
    value: Any = None
    confidence: float = 0.0
    source: str = "none"            # rules | heuristic | ai | ai+rules | conflict | excel | none
    evidence: Optional[str] = None
    note: Optional[str] = None

    @property
    def needs_review(self) -> bool:
        return self.value is not None and self.confidence < REVIEW_THRESHOLD


class ExtractionReport(BaseModel):
    filename: str
    method: str                     # "rules" | "rules+ai" | "ai-vision"
    fields: Dict[str, FieldResult]
    fine_print: List[str] = Field(default_factory=list)
    raw_text: str = ""
    ai_used: bool = False
    ai_error: Optional[str] = None

    @property
    def supplier_name(self) -> str:
        f = self.fields.get("supplier_name")
        return str(f.value) if f and f.value else "Unknown Supplier"

    def pending_review(self) -> List[str]:
        return [k for k, f in self.fields.items() if f.needs_review and k != "supplier_name"]

    def best_guess_quotation(self) -> SupplierQuotation:
        """Every extracted value regardless of confidence — used only for fraud screening, never for decisions."""
        return self.to_quotation({k: f.value for k, f in self.fields.items()
                                  if f.value is not None and k != "supplier_name"})

    def to_quotation(self, confirmed: Dict[str, Any]) -> SupplierQuotation:
        """Build the quotation the engine sees. Values below the confidence threshold are
        withheld (None) unless a human has confirmed them via `confirmed`."""
        data: Dict[str, Any] = {"supplier_name": self.supplier_name}
        for name, f in self.fields.items():
            if name == "supplier_name":
                continue
            if name in confirmed:
                data[name] = confirmed[name]
            elif f.value is not None and not f.needs_review:
                data[name] = f.value
        if data.get("gst_percent") is None:
            data.pop("gst_percent", None)
        if data.get("shipping_cost_total") is None:
            data.pop("shipping_cost_total", None)
        data["fine_print"] = list(self.fine_print)
        return SupplierQuotation(**data)


# ---------------------------------------------------------------------------
# Deterministic reader with evidence + confidence
# ---------------------------------------------------------------------------


def _evidence(text: str, m: re.Match) -> str:
    start = text.rfind("\n", 0, m.start()) + 1
    end = text.find("\n", m.end())
    end = len(text) if end == -1 else end
    snippet = text[start:end].strip()
    return snippet[:200]


def _try(patterns: List[re.Pattern], text: str, conf_primary: float, conf_fallback: float):
    for i, p in enumerate(patterns):
        m = p.search(text)
        if m:
            return m, (conf_primary if i == 0 else conf_fallback)
    return None, 0.0


WORD_NUMBERS = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7,
                "eight": 8, "nine": 9, "ten": 10, "twelve": 12, "fifteen": 15, "sixteen": 16,
                "twenty": 20, "twenty-four": 24, "thirty": 30, "thirty-two": 32, "sixty-four": 64}

HEURISTICS: Dict[str, List[Tuple[re.Pattern, Any, float]]] = {
    # (pattern, converter(match) -> value, confidence)
    "ram_gb": [
        (re.compile(r"(?:memory|ram)[^\n.]{0,25}?(\d+)\s*gigs?\b", re.I), lambda m: int(m.group(1)), 0.65),
        (re.compile(r"(\d+)\s*gigs?\b[^\n.]{0,15}(?:memory|ram)", re.I), lambda m: int(m.group(1)), 0.65),
    ],
    "storage_gb": [
        (re.compile(r"half\s+a\s+(?:terabyte|tb)", re.I), lambda m: 512, 0.55),
        (re.compile(r"(?:one|a)\s+(?:terabyte|tb)\b", re.I), lambda m: 1024, 0.6),
        (re.compile(r"(\d+)\s*gigs?\b[^\n.]{0,15}(?:ssd|nvme|storage|drive)", re.I), lambda m: int(m.group(1)), 0.6),
    ],
    "warranty_months": [
        (re.compile(r"(\d+)\s*(yrs?|years?|months?|mths?)\s*(?:on-?site\s*|carry-?in\s*)?(?:cover|warranty|guarantee)", re.I),
         lambda m: int(m.group(1)) * (12 if m.group(2).lower().startswith("y") else 1), 0.7),
        (re.compile(r"(?:cover|warranty|guarantee)[^\n.]{0,20}?(\d+)\s*(yrs?|years?|months?|mths?)", re.I),
         lambda m: int(m.group(1)) * (12 if m.group(2).lower().startswith("y") else 1), 0.7),
    ],
    "delivery_days": [
        (re.compile(r"(?:dispatch|deliver\w*|lead\s*time|supply)[^\n.]{0,30}?(?:a|one)\s+fortnight", re.I), lambda m: 14, 0.6),
        (re.compile(r"(?:dispatch|deliver\w*|lead\s*time|supply)[^\n.]{0,30}?(\d+)\s*weeks?", re.I), lambda m: int(m.group(1)) * 7, 0.65),
        (re.compile(r"(?:dispatch|deliver\w*|supply)[^\n.]{0,30}?(\d+)\s*(?:working\s*|calendar\s*)?days?", re.I), lambda m: int(m.group(1)), 0.7),
    ],
    "shipping_cost_total": [
        (re.compile(r"(?:freight|shipping|transport)[^\n.]{0,20}(?:extra|at\s+actuals?|additional|chargeable)", re.I),
         lambda m: None, 0.3),
    ],
}

GSTIN_RE = re.compile(r"\b(\d{2}[A-Z]{5}\d{4}[A-Z][1-9A-Z]Z[0-9A-Z])\b")
MSME_RE = re.compile(r"\b(?:MSME|Udyam|UDYAM-[A-Z]{2}-\d{2}-\d+|Micro\s*,?\s*Small)", re.I)
VALIDITY_RE = re.compile(r"(?:valid(?:ity)?|prices?\s+(?:are\s+)?(?:firm|valid))[^\n.]{0,25}?(\d+)\s*days", re.I)
PAYMENT_RE = re.compile(r"(?:payment|credit)[^\n.]{0,40}?(\d+)\s*days", re.I)

FINE_PRINT_RULES: List[Tuple[re.Pattern, str]] = [
    (re.compile(r"[^\n.]*batter(?:y|ies)[^\n.]*(?:only|exclud|not\s+covered|limited)[^\n.]*", re.I), "Battery coverage limited"),
    (re.compile(r"[^\n.]*(?:installation|imaging|setup|configuration)[^\n.]*(?:extra|additional|chargeable|charged)[^\n.]*", re.I), "Installation / setup charged extra"),
    (re.compile(r"[^\n.]*(?:freight|shipping|transport)[^\n.]*(?:extra|at\s+actuals?|additional|chargeable)[^\n.]*", re.I), "Freight not included"),
    (re.compile(r"[^\n.]*(?:valid|validity)[^\n.]*\b([1-9]|1[0-4])\s*days[^\n.]*", re.I), "Very short price validity"),
    (re.compile(r"[^\n.]*(?:advance|upfront)[^\n.]*(?:payment)?[^\n.]*", re.I), "Advance payment demanded"),
    (re.compile(r"[^\n.]*(?:refurbished|renewed|open[-\s]?box|grade\s*[ab]\b)[^\n.]*", re.I), "Refurbished / non-new goods"),
    (re.compile(r"[^\n.]*(?:ex-?works|ex-?warehouse)[^\n.]*", re.I), "Ex-works pricing (buyer bears logistics)"),
    (re.compile(r"[^\n.]*(?:subject\s+to\s+(?:change|availability|stock))[^\n.]*", re.I), "Price/stock subject to change"),
    (re.compile(r"[^\n.]*(?:carry-?in|return\s+to\s+base)[^\n.]*warranty[^\n.]*", re.I), "Carry-in (not on-site) warranty"),
]


def detect_fine_print(text: str) -> List[str]:
    found: List[str] = []
    for pattern, label in FINE_PRINT_RULES:
        for m in pattern.finditer(text):
            clause = re.sub(r"\s+", " ", m.group(0)).strip(" -•:;")
            if 8 <= len(clause) <= 220:
                item = f"{label}: “{clause}”"
                if item not in found:
                    found.append(item)
                break
    return found


def _sanity(field: str, value: Any, conf: float) -> Tuple[float, Optional[str]]:
    """Lower confidence for implausible values."""
    if value is None:
        return conf, None
    if field == "ram_gb" and value not in {4, 8, 12, 16, 24, 32, 48, 64, 96, 128}:
        return min(conf, 0.5), "Unusual RAM size"
    if field == "storage_gb" and value not in {128, 256, 512, 1000, 1024, 2048, 4096}:
        return min(conf, 0.55), "Unusual storage size"
    if field == "unit_price" and not (5_000 <= value <= 500_000):
        return min(conf, 0.45), "Price outside the plausible laptop range"
    if field == "warranty_months" and value > 60:
        return min(conf, 0.5), "Warranty longer than 5 years is unusual"
    if field == "delivery_days" and value > 120:
        return min(conf, 0.5), "Delivery over 120 days is unusual"
    if field == "gst_percent" and value not in {0, 5, 12, 18, 28}:
        return min(conf, 0.5), "Not a standard GST slab"
    return conf, None


def rules_read_text(text: str, filename: str) -> Dict[str, FieldResult]:
    out: Dict[str, FieldResult] = {}
    norm_text = "\n".join(normalise_casual_text(line) for line in text.splitlines())

    # supplier name
    m = rules._first_match(rules.SUPPLIER_PATTERNS, text)
    if m:
        out["supplier_name"] = FieldResult(field="supplier_name", value=rules._clean_supplier_name(m.group(1)),
                                           confidence=0.9, source="rules", evidence=_evidence(text, m))
    else:
        q = rules.parse_quotation_text(text, filename)
        out["supplier_name"] = FieldResult(field="supplier_name", value=q.supplier_name, confidence=0.75,
                                           source="rules", evidence="First line of document / file name")

    specs = [
        ("ram_gb", rules.RAM_PATTERNS, lambda m: int(m.group(1))),
        ("storage_gb", rules.STORAGE_PATTERNS,
         lambda m: int(m.group(1)) * (1024 if m.group(2).upper() == "TB" else 1)),
        ("unit_price", rules.PRICE_PATTERNS, lambda m: rules._clean_number(m.group(1))),
        ("warranty_months", rules.WARRANTY_PATTERNS,
         lambda m: int(m.group(1)) * (12 if m.group(2).lower().startswith("year") else 1)),
        ("delivery_days", rules.DELIVERY_PATTERNS, lambda m: int(m.group(1))),
        ("shipping_cost_total", rules.SHIPPING_PATTERNS, lambda m: rules._clean_number(m.group(1))),
        ("gst_percent", rules.GST_PATTERNS, lambda m: rules._clean_number(m.group(1))),
    ]
    for field, patterns, conv in specs:
        m, conf = _try(patterns, text, 0.95, 0.82)
        if m:
            value = conv(m)
            conf, note = _sanity(field, value, conf)
            out[field] = FieldResult(field=field, value=value, confidence=conf, source="rules",
                                     evidence=_evidence(text, m), note=note)
            continue
        # Messy wording: heuristics on the raw text, then the standard rules and heuristics on a
        # casual-phrase-normalised copy ("sixteen gigs", "two years cover"). Always below the review threshold.
        hit = None
        for pattern, conv_h, conf_h in HEURISTICS.get(field, []):
            hm = pattern.search(text)
            if hm:
                hit = (conv_h(hm), conf_h, _evidence(text, hm))
                break
        if hit is None:
            nm, _ = _try(patterns, norm_text, 0.75, 0.72)
            if nm:
                hit = (conv(nm), 0.75 if nm.re is patterns[0] else 0.72, _evidence(norm_text, nm))
        if hit is None:
            for pattern, conv_h, conf_h in HEURISTICS.get(field, []):
                hm = pattern.search(norm_text)
                if hm:
                    hit = (conv_h(hm), conf_h, _evidence(norm_text, hm))
                    break
        if hit is not None:
            value, conf_h, evidence = hit
            out[field] = FieldResult(
                field=field, value=value, confidence=conf_h if value is not None else 0.0, source="heuristic",
                evidence=evidence,
                note=("Non-standard wording — interpreted heuristically" if value is not None
                      else "Mentioned but no amount stated"),
            )
        else:
            out[field] = FieldResult(field=field, value=None, confidence=0.0, source="none")

    m = GSTIN_RE.search(text)
    out["gstin"] = FieldResult(field="gstin", value=m.group(1) if m else None, confidence=0.97 if m else 0.0,
                               source="rules" if m else "none", evidence=_evidence(text, m) if m else None)
    m = MSME_RE.search(text)
    out["is_msme"] = FieldResult(field="is_msme", value=True if m else None, confidence=0.9 if m else 0.0,
                                 source="rules" if m else "none", evidence=_evidence(text, m) if m else None)
    m = VALIDITY_RE.search(text)
    out["quote_validity_days"] = FieldResult(field="quote_validity_days", value=int(m.group(1)) if m else None,
                                             confidence=0.88 if m else 0.0, source="rules" if m else "none",
                                             evidence=_evidence(text, m) if m else None)
    m = PAYMENT_RE.search(text)
    out["payment_terms_days"] = FieldResult(field="payment_terms_days", value=int(m.group(1)) if m else None,
                                            confidence=0.85 if m else 0.0, source="rules" if m else "none",
                                            evidence=_evidence(text, m) if m else None)
    return out


def rules_read_excel(file_bytes: bytes, filename: str) -> Dict[str, FieldResult]:
    q = rules.parse_excel_quotation(file_bytes, filename)
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    evidence: Dict[str, str] = {}
    labelled: set = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = list(row)
            idx = next((i for i, c in enumerate(cells) if isinstance(c, str) and c.strip()), None)
            if idx is None:
                continue
            label = cells[idx]
            value = cells[idx + 1] if idx + 1 < len(cells) else None
            key = rules._match_excel_key(label)
            low = label.lower()
            if key is None:
                if "gstin" in low:
                    key = "gstin"
                elif "msme" in low or "udyam" in low:
                    key = "is_msme"
                elif "validity" in low:
                    key = "quote_validity_days"
                elif "payment" in low:
                    key = "payment_terms_days"
            if key and key not in evidence:
                evidence[key] = f"{label} | {'' if value is None else value}"
                labelled.add(key)
                if key == "gstin" and value:
                    q = q.model_copy(update={"gstin": str(value).strip()})
                if key == "is_msme" and value is not None:
                    q = q.model_copy(update={"is_msme": str(value).strip().lower() in {"yes", "y", "true", "1"}
                                             or "udyam" in str(value).lower()})
                if key == "quote_validity_days":
                    n = rules._excel_number(value)
                    q = q.model_copy(update={"quote_validity_days": int(n) if n is not None else None})
                if key == "payment_terms_days":
                    n = rules._excel_number(value)
                    q = q.model_copy(update={"payment_terms_days": int(n) if n is not None else None})

    out: Dict[str, FieldResult] = {}
    for field in ["supplier_name"] + NUMERIC_FIELDS + ["gstin", "is_msme"]:
        value = getattr(q, field, None)
        if field == "gst_percent" and "gst_percent" not in labelled:
            value = None
        if field == "shipping_cost_total" and "shipping_cost_total" not in labelled:
            value = None
        if value is None:
            out[field] = FieldResult(field=field, value=None, confidence=0.0, source="none",
                                     evidence=evidence.get(field),
                                     note="Labelled but left blank" if field in labelled else None)
        else:
            conf, note = _sanity(field, value, 0.97)
            out[field] = FieldResult(field=field, value=value, confidence=conf, source="excel",
                                     evidence=evidence.get(field), note=note)
    return out


# ---------------------------------------------------------------------------
# Claude reader
# ---------------------------------------------------------------------------


class AIField(BaseModel):
    value: Optional[str] = Field(None, description="The value exactly as a plain number or text, or null if not stated")
    confidence: float = Field(description="0.0-1.0 confidence that the value is correct and explicitly stated")
    evidence: Optional[str] = Field(None, description="Verbatim sentence from the document supporting the value")


class AIQuotation(BaseModel):
    supplier_name: AIField
    unit_price: AIField
    ram_gb: AIField
    storage_gb: AIField
    warranty_months: AIField
    delivery_days: AIField
    gst_percent: AIField
    shipping_cost_total: AIField
    gstin: AIField
    is_msme: AIField
    quote_validity_days: AIField
    payment_terms_days: AIField
    fine_print: List[str] = Field(default_factory=list,
                                  description="Clauses that add hidden cost or risk (exclusions, extras, short validity)")


AI_SYSTEM = """You extract structured data from supplier quotations for laptop procurement.
Rules:
- Report ONLY what the document explicitly states. If a value is not stated, return value=null and confidence 0.
  Never infer a default (e.g. do not assume a standard warranty).
- Normalise units: unit_price in INR per unit excluding GST; ram_gb and storage_gb in GB (1 TB = 1024 GB);
  warranty_months in months (1 year = 12); delivery_days in calendar days (1 week = 7, a fortnight = 14);
  shipping_cost_total in INR for the whole consignment; gst_percent as a number; is_msme as "true"/"false".
- If freight is "extra" or "at actuals" without an amount, shipping_cost_total is null and add it to fine_print.
- confidence reflects how explicit and unambiguous the statement is (vague wording => below 0.8).
- evidence must be copied verbatim from the document."""


def _coerce(field: str, raw: Optional[str]) -> Any:
    if raw is None:
        return None
    text = str(raw).strip()
    if text == "" or text.lower() in {"null", "none", "n/a", "not stated"}:
        return None
    if field == "supplier_name" or field == "gstin":
        return text
    if field == "is_msme":
        return text.lower() in {"true", "yes", "1"}
    num = rules._clean_number(re.sub(r"[^0-9.,\-]", "", text))
    if num is None:
        return None
    return int(round(num)) if field in INT_FIELDS else float(num)


def ai_read(text: str, file_bytes: bytes, filename: str) -> Optional[AIQuotation]:
    is_pdf = filename.lower().endswith(".pdf")
    if is_pdf and len(text.strip()) < 40:
        content = [llm.pdf_document_block(file_bytes),
                   llm.text_block("Extract the quotation fields from this (scanned) document.")]
    else:
        content = f"Quotation file: {filename}\n\n<document>\n{text}\n</document>\n\nExtract the quotation fields."
    return llm.structured_call(AI_SYSTEM, content, AIQuotation, effort="medium")


def reconcile(rule_fields: Dict[str, FieldResult], ai: AIQuotation) -> Dict[str, FieldResult]:
    merged: Dict[str, FieldResult] = {}
    for field in ["supplier_name"] + NUMERIC_FIELDS + ["gstin", "is_msme"]:
        r = rule_fields.get(field) or FieldResult(field=field)
        a_raw: AIField = getattr(ai, field)
        a_val = _coerce(field, a_raw.value)
        a_conf = max(0.0, min(1.0, a_raw.confidence)) if a_val is not None else 0.0
        if a_val is not None:
            a_conf, a_note = _sanity(field, a_val, a_conf)
        else:
            a_note = None

        def same(x, y) -> bool:
            if isinstance(x, (int, float)) and isinstance(y, (int, float)):
                return abs(float(x) - float(y)) <= max(0.5, 0.005 * abs(float(y)))
            return str(x).strip().lower() == str(y).strip().lower()

        if r.value is not None and a_val is not None:
            if same(r.value, a_val):
                merged[field] = FieldResult(field=field, value=r.value, confidence=min(0.99, max(r.confidence, a_conf) + 0.05),
                                            source="ai+rules", evidence=a_raw.evidence or r.evidence, note=r.note)
            else:
                pick_ai = a_conf >= r.confidence
                merged[field] = FieldResult(
                    field=field, value=a_val if pick_ai else r.value, confidence=0.5, source="conflict",
                    evidence=a_raw.evidence if pick_ai else r.evidence,
                    note=f"Rules read {r.value!r}, Claude read {a_val!r} — human must confirm",
                )
        elif a_val is not None:
            merged[field] = FieldResult(field=field, value=a_val, confidence=a_conf, source="ai",
                                        evidence=a_raw.evidence, note=a_note)
        elif r.value is not None:
            merged[field] = r.model_copy(update={"confidence": round(r.confidence * 0.85, 2),
                                                 "note": (r.note or "") + " Claude did not find this value."})
        else:
            merged[field] = FieldResult(field=field, value=None, confidence=0.0, source="none", note=r.note)
    return merged


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------


def extract_with_confidence(file_bytes: bytes, filename: str, use_ai: bool = True) -> ExtractionReport:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        raw_text = rules.extract_pdf_text(file_bytes)
        rule_fields = rules_read_text(raw_text, filename)
    elif ext in {".xlsx", ".xlsm"}:
        raw_text = rules.extract_excel_text(file_bytes)
        rule_fields = rules_read_excel(file_bytes, filename)
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Please upload .pdf or .xlsx quotations.")

    fine_print = detect_fine_print(raw_text)
    method, ai_used, ai_error, fields = "rules", False, None, rule_fields

    if use_ai and llm.llm_configured():
        ai = ai_read(raw_text, file_bytes, filename)
        if ai is not None:
            fields = reconcile(rule_fields, ai)
            ai_used = True
            method = "ai-vision" if ext == ".pdf" and len(raw_text.strip()) < 40 else "rules+ai"
            for clause in ai.fine_print:
                item = f"AI-flagged: “{clause.strip()}”"
                if item not in fine_print:
                    fine_print.append(item)
        else:
            ai_error = llm.last_error()

    if ext == ".pdf" and len(raw_text.strip()) < 40 and not ai_used:
        ai_error = ((ai_error + " ") if ai_error else "") + ("Document appears to be scanned (no text layer) — "
                                                              "please enter its values via the clarification loop.")

    return ExtractionReport(filename=filename, method=method, fields=fields, fine_print=fine_print,
                            raw_text=raw_text, ai_used=ai_used, ai_error=ai_error)


def mandatory_pending(report: ExtractionReport) -> List[str]:
    return [f for f in MANDATORY_FIELDS if f in report.pending_review()]
